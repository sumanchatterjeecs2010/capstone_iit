"""Build Report.docx and Report.pdf for the Multimodal Medical Assistant."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

OUT_DIR = Path(__file__).resolve().parent

TITLE = "Multimodal Medical Assistant for Image-Text Clinical Triage"
SUBTITLE = "HPPCS[04] — MedGemma 1.5 4B, Llama 3.2, LangGraph and Clinician Dashboard"

PARAS = [
    (
        "Abstract",
        "This project implements an educational Multimodal Medical Assistant (HPPCS[04]) "
        "that combines a medical vision LLM with a text LLM under LangGraph control. "
        "MedGemma 1.5 4B reads the uploaded medical image together with a clinical note; "
        "Llama 3.2 extracts entities, writes structured triage JSON, and produces clinician-facing "
        "text. The default entry point is python main.py, which opens a browser dashboard where "
        "clinicians upload de-identified note and image files, review triage and image findings, "
        "and ask interactive follow-up questions. An optional batch mode (python main.py --batch) "
        "processes five public teaching cases stored in sample_data/. Privacy scrubbing, entity "
        "normalization, and educational reference links are included. This is a prototype, not a "
        "diagnostic device.",
    ),
    (
        "1. Introduction",
        "Multimodal clinical support requires both image understanding and structured reasoning "
        "over text. A single undifferentiated script makes errors hard to trace—for example, "
        "asking a text-only model to invent imaging signs. Earlier versions used separate encoders "
        "(OpenCV, CLIP, BioBERT) plus one generator. The current design uses two generative LLMs "
        "with explicit roles and a LangGraph orchestrator. OpenCV, CLIP, BioBERT and LangChain "
        "LCEL chains were removed.",
    ),
    (
        "2. Problem Statement",
        "Given a medical image and a short clinical note, produce an educational assistant response "
        "that correlates image findings with symptoms, assigns triage (emergency, urgent, soon or "
        "routine), suggests next steps, and supports follow-up questioning. Clinicians must be "
        "able to upload their own files through a simple interface without mixing them with demo data.",
    ),
    (
        "3. Objectives",
        None,
    ),
    (
        "4. Two-LLM Architecture",
        "MedGemma 1.5 4B (medgemma:4b on Ollama) is a Gemma-3-based multimodal model for medical "
        "image and text. It receives the de-identified image and note and returns structured visual "
        "findings. Llama 3.2 3B (or 1B on low-RAM machines) is the text generator: entity extraction, "
        "triage JSON, initial summary, and follow-up chat replies. The models run sequentially through "
        "Ollama so both weight files need not stay in RAM at once.",
    ),
    (
        "5. LangGraph Workflow",
        "LangGraph is the orchestrator, not a medical model. A linear four-node graph shares typed "
        "state: (1) medgemma_analyze → visual_analysis; (2) llama_entities → extracted_entities; "
        "(3) llama_reason → clinical_reasoning with triage and recommendations; (4) llama_converse → "
        "one clinician-facing summary turn. Interactive follow-ups after upload use Llama text-only "
        "via POST /chat, reusing stored findings without re-running MedGemma. JSON parsing and retry "
        "logic live in ollama_client.py.",
    ),
    (
        "6. Ingestion, Privacy and NLP",
        "ingestion.py accepts clinical notes (.txt/.md) and images (JPEG, PNG, DICOM). privacy.py "
        "de-identifies text with regex redaction (names, emails, phones, MRNs, dates) and strips "
        "JPEG EXIF, PNG text chunks, or DICOM PHI tags before storage. Original uploads are never "
        "written to disk. entity_normalizer.py maps Llama output to local educational codes for "
        "conditions, medications, vitals and requested studies. clinical_references.py attaches "
        "public WHO/NIH guideline links keyed by condition keywords.",
    ),
    (
        "7. Clinician Interface",
        "upload_app.py (FastAPI) serves the dashboard at http://127.0.0.1:8000/. The left panel "
        "shows the processed image, triage badge, visual findings, impression, differential and "
        "reference links. The right panel is an interactive chat: clinicians type follow-up "
        "questions or click suggested chips; triage may update when new red-flag information "
        "appears. Each session is stored under uploads/processed/<session_id>/conversation.json. "
        "Teaching demo files live separately in sample_data/ and are not required for user uploads.",
    ),
    (
        "8. Dataset and Evaluation",
        "dataset_builder.py downloads five public teaching images from Wikimedia/NIH and writes "
        "matching fictional notes into sample_data/. python main.py --batch runs the pipeline on "
        "these cases and writes conversation_01.json … conversation_05.json plus "
        "evaluation_summary.json with keyword recall and triage match against known teaching labels. "
        "These metrics are educational only, not clinical validation. MedGemma vision is slow on "
        "CPU; a GPU (e.g. Google Colab T4) is recommended for first analysis.",
    ),
    (
        "9. System Modules",
        "main.py — entry point (dashboard default, --batch demo, CLI --text/--image). "
        "graph_pipeline.py — LangGraph nodes and prompts. medical_assistant.py — case packaging, "
        "session load/save, follow-up handler. ollama_client.py — Ollama HTTP wrapper. "
        "upload_app.py — dashboard and chat API. Removed: image_analyzer.py, multimodal_encoder.py, "
        "langchain_orchestrator.py.",
    ),
    (
        "10. Limitations and Ethics",
        "The assistant must not be used for real patient care. Small LLMs hallucinate. Evaluation "
        "is keyword-based on teaching cases, not a reader study. De-identification is "
        "HIPAA-inspired but not certified. Users are responsible for lawful use of uploaded content.",
    ),
    (
        "11. Conclusion and Future Work",
        "The project meets HPPCS[04] with two LLMs, LangGraph orchestration, upload ingestion, "
        "de-identification, entity normalization, and a clinician dashboard with interactive chat. "
        "Future work: clinician-rated conversation quality, optional SNOMED/RxNorm coding, and "
        "evaluation on de-identified real studies with appropriate governance.",
    ),
    (
        "12. References",
        "[1] HAAI++ capstone instructions and HPPCS[04] brief. "
        "[2] A. Grattafiori et al., The Llama 3 herd of models, arXiv:2407.21783. "
        "[3] Google Health AI Developer Foundations, MedGemma model card, "
        "https://developers.google.com/health-ai-developer-foundations/medgemma/model-card "
        "[4] LangGraph documentation, https://langchain-ai.github.io/langgraph/ "
        "[5] Ollama MedGemma, https://ollama.com/library/medgemma "
        "[6] Ollama Llama 3.2, https://ollama.com/library/llama3.2 "
        "[7] Wikimedia Commons and NIH public teaching images (see sample_data/IMAGE_SOURCES.txt).",
    ),
    (
        "Acknowledgement",
        "Thanks to HAAI++ faculty and the MedGemma, Llama, LangGraph and Ollama communities. "
        "Teaching images are public; accompanying notes are fictional.",
    ),
]

OBJECTIVES = [
    "Provide a clinician dashboard to upload note + image and receive triage guidance.",
    "Use two LLMs: MedGemma 1.5 4B (vision) and Llama 3.2 (text).",
    "Orchestrate with LangGraph: analyze → entities → triage → summary.",
    "De-identify uploads before storage or model calls.",
    "Normalize clinical entities and attach educational reference links.",
    "Support interactive follow-up chat with adaptive triage updates.",
    "Optional batch evaluation on five teaching cases in sample_data/.",
]


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    title = doc.add_heading(TITLE, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(SUBTITLE)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in sub.runs:
        run.italic = True

    for heading, body in PARAS:
        doc.add_heading(heading, 1)
        if heading.startswith("3. Objectives"):
            for item in OBJECTIVES:
                doc.add_paragraph(item, style="List Bullet")
            continue
        p = doc.add_paragraph(body)
        p.paragraph_format.space_after = Pt(8)

    path = OUT_DIR / "Report.docx"
    doc.save(path)
    return path


def build_pdf():
    path = OUT_DIR / "Report.pdf"
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="CoverTitle",
            parent=styles["Title"],
            fontSize=16,
            leading=20,
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CoverSub",
            parent=styles["Normal"],
            fontSize=11,
            leading=14,
            alignment=1,
            spaceAfter=18,
            italic=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="BodyJust",
            parent=styles["BodyText"],
            fontSize=11,
            leading=15,
            spaceAfter=10,
        )
    )
    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title=TITLE,
        author="Capstone HPPCS[04]",
    )
    story = [
        Paragraph(TITLE, styles["CoverTitle"]),
        Paragraph(SUBTITLE, styles["CoverSub"]),
    ]
    for heading, body in PARAS:
        story.append(Paragraph(heading, styles["Heading1"]))
        if heading.startswith("3. Objectives"):
            story.append(
                ListFlowable(
                    [ListItem(Paragraph(item, styles["BodyJust"])) for item in OBJECTIVES],
                    bulletType="bullet",
                    leftIndent=18,
                )
            )
            story.append(Spacer(1, 8))
            continue
        story.append(Paragraph(body, styles["BodyJust"]))
    doc.build(story)
    return path


if __name__ == "__main__":
    print("Wrote", build_docx())
    print("Wrote", build_pdf())
