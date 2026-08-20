"""Build Report/Report.docx and Report/Report.pdf (run from project root).

Target length: 3 pages (Times New Roman 12, 1.5 line spacing).
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPORT_DIR = Path(__file__).resolve().parent / "Report"

TITLE = "Multimodal Medical Assistant for Image-Text Clinical Triage"

ABSTRACT = (
    "This project implements a clinician-facing multimodal assistant that combines a radiology "
    "scan or histopathology slide with a structured clinical note to produce triage guidance, "
    "image–text correlation, evidence-linked summaries, and interactive follow-up chat. The "
    "clinical aim is to reduce missed cross-modal cues by forcing vision and text models to share "
    "one case record. Compact local models are used deliberately: MedGemma 4B for vision and "
    "Llama 3.2 (3B, with 1B fallback) for text reasoning, chosen for memory efficiency on a "
    "typical 16 GB CPU laptop or Google Colab (CPU or free T4 GPU). Ollama auto-adapts to CPU or "
    "GPU; the app unloads the vision model after analysis so both LLMs share limited RAM/VRAM. "
    "Modalities are restricted to radiology and pathology with auto domain detection. LangGraph "
    "sequences analysis and follow-up re-triage. The CLI takes --text and --image paths, prints "
    "triage summaries, continues into follow-up chat, and writes conversation.json in the Codebase "
    "directory; optional --tui offers the same pipeline with interactive path prompts. Label-free "
    "evaluation metrics are printed after every run."
)

OBJECTIVES = [
    "Provide a CLI (and optional TUI) so clinicians pass a note and image path and receive triage summaries.",
    "Return image-referenced findings and contextual explanations from the clinical note.",
    "Support adaptive follow-up with Llama re-triage of impression, differential, recommendations, and next questions.",
    "Auto-detect radiology versus pathology and reject out-of-scope images.",
    "De-identify text and images before any model call.",
    "Emit evaluation metrics (correlation, robustness, explanation quality, satisfaction) on every run.",
    "Write the conversation output as conversation.json in the Codebase directory.",
]


def _set_run_font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def _add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    _set_run_font(run, size=10)


def _para(doc, text, *, size=12, bold=False, italic=False, center=False, space_after=4):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    _set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def _heading(doc, text):
    return _para(doc, text, size=12, bold=True, space_after=3)


def _bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(1)
    pf.space_before = Pt(0)
    if p.runs:
        p.runs[0].text = text
        _set_run_font(p.runs[0], size=12)
    else:
        run = p.add_run(text)
        _set_run_font(run, size=12)
    return p


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number(fp)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    _para(doc, TITLE, size=15, bold=True, center=True, space_after=6)
    _para(doc, "HPPCS[04] Capstone Report", size=12, italic=True, center=True, space_after=8)

    _heading(doc, "Abstract")
    _para(doc, ABSTRACT, space_after=4)

    _heading(doc, "1. Introduction")
    _bullet(
        doc,
        "Context: Imaging and notes are usually read separately; this system binds them in one "
        "LangGraph state so triage uses visual findings and note context together.",
    )
    _bullet(
        doc,
        "Motivation: Local MedGemma + Llama on Ollama keep PHI on-premises and match HPPCS[04].",
    )

    _heading(doc, "2. Problem Statement")
    _para(
        doc,
        "Given one radiology or pathology image and one clinical note, the system must detect "
        "domain, correlate image findings with the note, assign triage (emergency / urgent / soon / "
        "routine), answer follow-up questions with updated recommendations, and report evaluation "
        "metrics. Out-of-scope photographs (dermatology, fundus, clinical snapshots) are rejected.",
        space_after=4,
    )

    _heading(doc, "3. Objectives")
    for item in OBJECTIVES:
        _bullet(doc, item)

    _heading(doc, "4. Methodology")
    _bullet(
        doc,
        "Stack: Python 3.12, LangGraph, Ollama, MedGemma 4B, Llama 3.2 3B, runtime_profile.py.",
    )
    _bullet(
        doc,
        "Workflow: de-identify → MedGemma (domain + findings) → Llama entities → Llama triage JSON → "
        "summary → follow-up re-triage; metrics via evaluation.py.",
    )
    _para(
        doc,
        "Pipeline: Note+Image → Privacy → MedGemma → Llama → conversation.json → follow-up chat.",
        italic=True,
        space_after=4,
    )

    _heading(doc, "5. System Design / Implementation")
    _para(
        doc,
        "Architecture (layers): Presentation (main.py CLI and optional tui_app.py) → Application "
        "(medical_assistant.py) → Orchestration (graph_pipeline.py / LangGraph) → Models "
        "(ollama_client.py with runtime_profile.py for CPU/GPU) → Supporting modules "
        "(ingestion.py, privacy.py, evaluation.py, paths.py). After each run, "
        "medical_assistant.py writes conversation.json in the Codebase directory.",
        space_after=2,
    )
    _bullet(doc, "main.py — CLI with --text and --image; follow-up chat after analysis; optional --tui; --device auto|cpu|gpu.")
    _bullet(doc, "runtime_profile.py — detects NVIDIA GPU vs CPU; tunes keep-alive, timeouts, and unload policy.")
    _bullet(doc, "graph_pipeline.py — four-node LangGraph; follow-up re-triage grounded in image and note.")
    _bullet(doc, "evaluation.py — correlation, robustness, explanation quality, and satisfaction metrics.")
    _bullet(doc, "tui_app.py — shared terminal summary and follow-up chat loop used by CLI and --tui.")
    _bullet(doc, "execution.txt — setup and run instructions for the submission.")

    _heading(doc, "6. Observations and Results")
    _para(
        doc,
        "Evaluation is performed during a live CLI or TUI run: pass note and image paths, wait for "
        "analysis, continue follow-up chat at the You> prompt, then read printed metrics and "
        "Codebase/conversation.json. Faculty may supply their own note and image at demo time.",
        space_after=3,
    )
    _para(
        doc,
        "On representative radiology and pathology cases, MedGemma distinguishes domains; Llama "
        "returns structured triage (emergency / urgent / soon / routine) with impression, "
        "differential, and recommendations; image–note correlation links findings to the note; "
        "and follow-up questions trigger re-triage. Quantitative metrics from evaluation.py:",
        space_after=2,
    )
    _bullet(doc, "Cross-modal correlation — image–note narrative presence and token overlap.")
    _bullet(doc, "Robustness — valid domain detection, de-identified processing, and field completeness.")
    _bullet(doc, "Explanation quality — checklist of clinical explanation fields (mapped to 1–5 Likert).")
    _bullet(doc, "User satisfaction — interface completeness proxy from triage, findings, chat, and prompts.")
    _para(
        doc,
        "Demo command: python main.py --text <note.txt> --image <scan.jpg> from Codebase/.",
        space_after=4,
    )

    _heading(doc, "7. Conclusion and Future Work")
    _bullet(
        doc,
        "Summary of contributions: dual-LLM multimodal triage with auto domain detection, privacy "
        "gate, CLI (and optional TUI), adaptive follow-up, and evaluation metric outputs.",
    )
    _bullet(
        doc,
        "Possible extensions: larger GPU serving when hardware allows, and reader-study Likert "
        "collection from clinicians.",
    )

    _heading(doc, "8. References")
    _bullet(doc, "HPPCS[04] capstone brief.")
    _bullet(doc, "Google Health AI Developer Foundations, MedGemma model card.")
    _bullet(doc, "Meta Llama 3.2 model card; Ollama library (medgemma, llama3.2).")
    _bullet(doc, "LangGraph documentation; Wikimedia Commons teaching images in sample_data/.")

    _heading(doc, "Acknowledgement")
    _para(
        doc,
        "Thanks to HAAI++ faculty for the HPPCS[04] brief and to the MedGemma, Llama, LangGraph, "
        "and Ollama communities. Teaching images are public; accompanying notes are fictional.",
        space_after=2,
    )

    out = Path(REPORT_DIR) / "Report.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(out)
    except PermissionError:
        out = Path(REPORT_DIR) / "Report_new.docx"
        doc.save(out)
        print("Report.docx is open — wrote", out, "instead. Close Word and rerun build_report.py.")
    return out


def docx_to_pdf(docx_path):
    docx_path = Path(docx_path)
    pdf_path = docx_path.with_suffix(".pdf")
    try:
        from docx2pdf import convert

        convert(str(docx_path), str(pdf_path))
        return pdf_path
    except Exception as exc:
        print("docx2pdf failed ({}); trying Word COM…".format(exc))
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    doc = word.Documents.Open(str(docx_path.resolve()))
    doc.SaveAs(str(pdf_path.resolve()), FileFormat=17)
    doc.Close()
    word.Quit()
    return pdf_path


if __name__ == "__main__":
    docx = build_docx()
    print("Wrote", docx)
    print("Wrote", docx_to_pdf(docx))
